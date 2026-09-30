#!/usr/bin/env python3
"""
Dump TWG Comp.xlsx into JSON for import and for the historical scoring replay.

The sheets are not uniform: the header row moves (row 23 in most seasons, 26 in
2018/19), the competition set changes every year, and the tournament sheets use
a completely different layout. So we locate the header by content rather than by
a fixed offset, and emit whatever rows we find rather than assuming a schema.

Usage: python3 scripts/dump-workbook.py <path-to-xlsx> > scripts/workbook.json
"""
import json
import sys

import openpyxl


def cell(ws, row, col):
    v = ws.cell(row=row, column=col).value
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def dump_season(ws):
    """Season sheets: label in col A, then (prediction, points) pairs per player."""
    header_row = None
    for r in range(1, min(ws.max_row, 40) + 1):
        if (cell(ws, r, 1) or "").lower() == "country":
            header_row = r
            break
    if header_row is None:
        return None

    player_row = header_row - 1
    players = []
    for col in (2, 4, 6):
        name = cell(ws, player_row, col)
        if name:
            players.append({"handle": name, "predCol": col, "ptsCol": col + 1})

    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        label = cell(ws, r, 1)
        if not label:
            continue
        if label.lower() in ("total", "honour board", "season:"):
            break
        entry = {
            "row": r,
            "label": label,
            "picks": {},
            # Column H carries the true Golden Boot tally on "Golden Boot #" rows.
            # It is the only actual result the workbook ever recorded.
            "actual": cell(ws, r, 8),
        }
        for p in players:
            entry["picks"][p["handle"]] = {
                "value": cell(ws, r, p["predCol"]),
                "points": cell(ws, r, p["ptsCol"]),
            }
        rows.append(entry)

    totals = {}
    for r in range(header_row + 1, ws.max_row + 1):
        if (cell(ws, r, 1) or "").lower() == "total":
            for p in players:
                totals[p["handle"]] = cell(ws, r, p["ptsCol"])
            break

    return {
        "players": [p["handle"] for p in players],
        "rows": rows,
        "sheetTotals": totals,
    }


def dump_honour(ws):
    out = []
    for r in range(1, ws.max_row + 1):
        label = cell(ws, r, 1)
        if not label or "/" not in label:
            continue
        out.append(
            {
                "season": label,
                "seasonWinner": cell(ws, r, 2),
                "tournament": cell(ws, r, 3),
                "tournamentWinner": cell(ws, r, 4),
            }
        )
    return out


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else "TWG Comp.xlsx"
    wb = openpyxl.load_workbook(path, data_only=True)

    out = {"seasons": {}, "honour": [], "skipped": []}
    for ws in wb.worksheets:
        title = ws.title
        if title == "Honour Board":
            out["honour"] = dump_honour(ws)
            continue
        data = dump_season(ws)
        if data is None:
            # Tournament sheets (WC/Euro) have no "Country" header row. Out of
            # scope for v1 — recorded so the gap is visible, not silent.
            out["skipped"].append(title)
            continue
        out["seasons"][title] = data

    json.dump(out, sys.stdout, indent=1)


if __name__ == "__main__":
    main()
